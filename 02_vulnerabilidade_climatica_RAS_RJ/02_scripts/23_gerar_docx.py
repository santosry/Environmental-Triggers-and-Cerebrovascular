#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Converte os arquivos Markdown do manuscrito para .docx (Word), atendendo ao
edital (manuscrito e folha de rosto em arquivos Word separados).

Suporta: titulos (#..####), paragrafos, negrito (**), italico (*), codigo inline,
listas (- / 1.), citacoes (>), tabelas (| ... |) e regras horizontais (---).
As citacoes ja estao em sobrescrito Unicode, preservadas no .docx.
"""
import os
import re
import sys

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH


def add_runs(par, texto):
    """Adiciona texto com **negrito**, *italico*, `codigo` e <sup>sobrescrito</sup>."""
    tokens = re.split(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`|<sup>.+?</sup>)", texto)
    for tk in tokens:
        if not tk:
            continue
        if tk.startswith("<sup>") and tk.endswith("</sup>"):
            r = par.add_run(tk[5:-6]); r.font.superscript = True
        elif tk.startswith("**") and tk.endswith("**"):
            r = par.add_run(tk[2:-2]); r.bold = True
        elif tk.startswith("*") and tk.endswith("*") and len(tk) > 2:
            r = par.add_run(tk[1:-1]); r.italic = True
        elif tk.startswith("`") and tk.endswith("`"):
            r = par.add_run(tk[1:-1]); r.font.name = "Consolas"
        else:
            par.add_run(tk)


def converter(md_path, docx_path):
    doc = Document()
    estilo = doc.styles["Normal"]
    estilo.font.name = "Times New Roman"
    estilo.font.size = Pt(12)
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Cm(2.5)

    linhas = open(md_path, encoding="utf-8").read().splitlines()
    i = 0
    while i < len(linhas):
        ln = linhas[i]
        s = ln.strip()

        if s in ("---", "***", "___"):
            i += 1
            continue
        if not s:
            i += 1
            continue

        # tabelas
        if s.startswith("|") and s.endswith("|"):
            bloco = []
            while i < len(linhas) and linhas[i].strip().startswith("|"):
                bloco.append(linhas[i].strip())
                i += 1
            rows = []
            for b in bloco:
                if re.match(r"^\|[\s:\-|]+\|$", b):
                    continue
                cells = [c.strip() for c in b.strip("|").split("|")]
                rows.append(cells)
            if rows:
                ncol = max(len(r) for r in rows)
                t = doc.add_table(rows=0, cols=ncol)
                t.style = "Table Grid"
                for ri, r in enumerate(rows):
                    cres = t.add_row().cells
                    for ci in range(ncol):
                        txt = r[ci] if ci < len(r) else ""
                        p = cres[ci].paragraphs[0]
                        add_runs(p, txt)
                        for run in p.runs:
                            run.font.size = Pt(10)
                            if ri == 0:
                                run.bold = True
                doc.add_paragraph("")
            continue

        # titulos
        m = re.match(r"^(#{1,6})\s+(.*)$", s)
        if m:
            nivel = len(m.group(1))
            doc.add_heading(m.group(2).replace("*", ""), level=min(nivel, 4))
            i += 1
            continue

        # citacao
        if s.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1.0)
            add_runs(p, s.lstrip("> ").strip())
            for r in p.runs:
                r.italic = True
                r.font.size = Pt(10.5)
            i += 1
            continue

        # listas
        if re.match(r"^[-*]\s+", s):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, re.sub(r"^[-*]\s+", "", s))
            i += 1
            continue
        if re.match(r"^\d+\.\s+", s):
            p = doc.add_paragraph(style="List Number")
            add_runs(p, re.sub(r"^\d+\.\s+", "", s))
            i += 1
            continue

        # paragrafo
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        if s.startswith("**Título") or s.startswith("**Palavras") or s.startswith("**Keywords"):
            add_runs(p, s)
        else:
            add_runs(p, s)
        i += 1

    doc.save(docx_path)
    return docx_path


def main():
    base = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    mds = ["manuscrito.md", "folha_de_rosto.md", "carta_apresentacao.md",
           "declaracao_etica.md", "termo_autoria_responsabilidade.md"]
    for md in mds:
        src = os.path.join(base, "08_manuscrito", md)
        if not os.path.exists(src):
            print("ausente:", src); continue
        dst = src[:-3] + ".docx"
        converter(src, dst)
        print("ok", os.path.basename(dst))


if __name__ == "__main__":
    main()
